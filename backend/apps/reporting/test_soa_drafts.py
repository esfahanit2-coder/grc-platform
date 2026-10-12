from io import BytesIO
from datetime import timedelta

from django.utils import timezone

from docx import Document as Docx
from django.contrib.auth import get_user_model
from rest_framework.test import APITestCase
from rest_framework_simplejwt.tokens import RefreshToken

from apps.controls.models import Control, ControlImplementation, ControlRequirement
from apps.frameworks.models import Framework, FrameworkVersion, Requirement
from apps.identity.models import UserRoleScope
from apps.identity.services import bootstrap_tenant_rbac
from apps.organizations.models import OrganizationUnit
from apps.risks.models import Risk, RiskControl, RiskTreatment
from apps.tenancy.models import Tenant, TenantMembership
from apps.evidence.models import Evidence, EvidenceLink
from apps.assessments.models import Assessment

from .models import SoAControlDecision, SoARegister


class SoADraftControlRegisterTests(APITestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user(username="soa-admin-test", password="StrongPassword123!")
        self.viewer = User.objects.create_user(username="soa-viewer-test", password="StrongPassword123!")
        self.tenant = Tenant.objects.create(name="SoA Tenant One", code="soa-tenant-one")
        TenantMembership.objects.create(tenant=self.tenant, user=self.admin, role_code="admin")
        TenantMembership.objects.create(tenant=self.tenant, user=self.viewer, role_code="viewer")
        roles = bootstrap_tenant_rbac(self.tenant, admin_user=self.admin)

        self.hq = OrganizationUnit.objects.create(
            tenant=self.tenant, unit_type=OrganizationUnit.UnitType.COMPANY, code="HQ", name="HQ"
        )
        self.ops = OrganizationUnit.objects.create(
            tenant=self.tenant, unit_type=OrganizationUnit.UnitType.BUSINESS_UNIT, code="OPS", name="OPS"
        )
        UserRoleScope.objects.create(
            user=self.viewer, tenant=self.tenant, role=roles["viewer"], organization_unit=self.hq
        )
        framework = Framework.objects.create(
            tenant=self.tenant, code="authorized-source-synthetic", name="Synthetic local framework"
        )
        self.version = FrameworkVersion.objects.create(
            framework=framework, version_code="synthetic-locked", is_locked=True
        )
        self.unlocked_version = FrameworkVersion.objects.create(
            framework=framework, version_code="synthetic-unlocked", is_locked=False
        )
        self.control = Control.objects.create(
            tenant=self.tenant, code="syn-ctrl-01", title="Synthetic control 01",
            status=Control.Status.ACTIVE
        )
        self.second_control = Control.objects.create(
            tenant=self.tenant, code="syn-ctrl-02", title="Synthetic control 02",
            status=Control.Status.ACTIVE
        )
        self.implementation = ControlImplementation.objects.create(
            tenant=self.tenant, organization_unit=self.hq, control=self.control, owner=self.admin
        )
        self.foreign_implementation_scope = ControlImplementation.objects.create(
            tenant=self.tenant, organization_unit=self.ops, control=self.second_control, owner=self.admin
        )
        self.risk = Risk.objects.create(
            tenant=self.tenant, organization_unit=self.hq, code="SYN-R-01",
            title="Synthetic source scope risk", owner=self.admin
        )
        self.treatment = RiskTreatment.objects.create(
            risk=self.risk, strategy=RiskTreatment.Strategy.REDUCE, owner=self.admin
        )
        RiskControl.objects.create(
            risk=self.risk, control_implementation=self.implementation,
            relationship_type=RiskControl.RelationshipType.TREATMENT,
        )
        self.foreign_tenant = Tenant.objects.create(name="SoA Tenant Two", code="soa-tenant-two")
        self.foreign_unit = OrganizationUnit.objects.create(
            tenant=self.foreign_tenant, unit_type=OrganizationUnit.UnitType.COMPANY,
            code="HQ", name="HQ Two"
        )
        self.foreign_control = Control.objects.create(
            tenant=self.foreign_tenant, code="private-other", title="Foreign tenant control"
        )

    def auth(self, user, tenant=None):
        token = str(RefreshToken.for_user(user).access_token)
        self.client.credentials(
            HTTP_AUTHORIZATION=f"Bearer {token}",
            HTTP_X_TENANT_ID=str((tenant or self.tenant).id),
        )

    def make_register(self, unit=None, revision=1):
        self.auth(self.admin)
        response = self.client.post("/api/v1/soa/registers/", {
            "organization_unit": str((unit or self.hq).id),
            "framework_version": str(self.version.id),
            "title": "Unapproved synthetic SoA decisions",
            "revision": revision,
        }, format="json")
        self.assertEqual(response.status_code, 201, response.content)
        return response.json()

    def make_decision(self, register, **overrides):
        self.auth(self.admin)
        payload = {
            "control": str(self.control.id),
            "reference_kind": "annex_a",
            "source_identifier": "A.5.1",
            "decision": "included",
            "rationale": "Example human-provided rationale, not ISO content",
            "control_implementation": str(self.implementation.id),
            "risk_treatment": str(self.treatment.id),
        }
        payload.update(overrides)
        response = self.client.post(
            f"/api/v1/soa/registers/{register['id']}/decisions/", payload, format="json"
        )
        return response

    def test_create_and_retrieve_draft_with_selected_control_and_risk(self):
        register = self.make_register()
        self.assertEqual(register["status"], "draft")
        self.assertFalse(register["is_formal_statement_of_applicability"])
        response = self.make_decision(register)
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(response.json()["risk_treatment"], str(self.treatment.id))
        self.assertEqual(response.json()["approval_status"], "draft_unapproved")
        self.assertEqual(SoAControlDecision.objects.filter(register_id=register["id"]).count(), 1)
        self.auth(self.viewer)
        preview = self.client.get(f"/api/v1/soa/registers/{register['id']}/")
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(len(preview.json()["decisions"]), 1)
        self.assertEqual(preview.json()["notice"].split(" — ")[0], "DRAFT / NOT APPROVED")

    def test_scoped_viewer_can_only_read_assigned_unit_and_cannot_write(self):
        good = self.make_register(self.hq)
        bad = self.make_register(self.ops)
        self.auth(self.viewer)
        listing = self.client.get("/api/v1/soa/registers/")
        self.assertEqual(listing.status_code, 200)
        self.assertEqual([x["id"] for x in listing.json()["results"]], [good["id"]])
        self.assertEqual(self.client.get(f"/api/v1/soa/registers/{bad['id']}/").status_code, 404)
        forbidden = self.client.post("/api/v1/soa/registers/", {
            "organization_unit": str(self.hq.id),
            "framework_version": str(self.version.id),
            "title": "not allowed",
        }, format="json")
        self.assertEqual(forbidden.status_code, 403)
        forbidden_decision = self.client.post(
            f"/api/v1/soa/registers/{good['id']}/decisions/",
            {"control": str(self.control.id)}, format="json"
        )
        self.assertEqual(forbidden_decision.status_code, 403)

    def test_wrong_tenant_and_unknown_controls_are_never_accepted(self):
        register = self.make_register()
        wrong = self.make_decision(register, control=str(self.foreign_control.id))
        self.assertEqual(wrong.status_code, 400)
        self.assertEqual(register["id"], str(SoARegister.objects.get(pk=register["id"]).id))
        self.assertEqual(SoAControlDecision.objects.count(), 0)
        self.auth(self.admin, tenant=self.foreign_tenant)
        other = self.client.get(f"/api/v1/soa/registers/{register['id']}/")
        self.assertIn(other.status_code, (403, 404))
        self.assertNotIn("Unapproved synthetic SoA decisions", other.content.decode())

    def test_unlocked_framework_source_is_rejected_and_approval_cannot_be_posted(self):
        self.auth(self.admin)
        unprotected = self.client.post("/api/v1/soa/registers/", {
            "organization_unit": str(self.hq.id),
            "framework_version": str(self.unlocked_version.id),
            "title": "unsafe source",
        }, format="json")
        self.assertEqual(unprotected.status_code, 400)
        forced = self.client.post("/api/v1/soa/registers/", {
            "organization_unit": str(self.hq.id),
            "framework_version": str(self.version.id),
            "title": "force approval", "status": "approved",
        }, format="json")
        self.assertEqual(forced.status_code, 400)
        self.assertEqual(SoARegister.objects.count(), 0)

    def test_inclusion_or_exclusion_requires_rationale_and_links_share_scope(self):
        register = self.make_register()
        no_reason = self.make_decision(register, rationale="")
        self.assertEqual(no_reason.status_code, 400)
        bad_scope = self.make_decision(
            register, control=str(self.second_control.id),
            control_implementation=str(self.foreign_implementation_scope.id)
        )
        self.assertEqual(bad_scope.status_code, 400)
        excluded_with_treatment = self.make_decision(register, decision="excluded")
        self.assertEqual(excluded_with_treatment.status_code, 400)
        self.assertEqual(SoAControlDecision.objects.count(), 0)

    def test_duplicate_and_foreign_scope_treatment_are_rejected(self):
        register = self.make_register()
        first = self.make_decision(register)
        self.assertEqual(first.status_code, 201, first.content)
        duplicate = self.make_decision(register)
        self.assertEqual(duplicate.status_code, 400)
        other_risk = Risk.objects.create(
            tenant=self.tenant, organization_unit=self.ops, code="SYN-R-02",
            title="Different scope", owner=self.admin
        )
        other_treatment = RiskTreatment.objects.create(
            risk=other_risk, strategy=RiskTreatment.Strategy.REDUCE, owner=self.admin
        )
        foreign_scope = self.make_decision(
            register, control=str(self.second_control.id), source_identifier="A.5.2",
            control_implementation=None, risk_treatment=str(other_treatment.id),
        )
        self.assertEqual(foreign_scope.status_code, 400)
        self.assertEqual(SoAControlDecision.objects.count(), 1)

    def test_same_scope_treatment_without_risk_control_link_is_rejected(self):
        register = self.make_register()
        unrelated_risk = Risk.objects.create(
            tenant=self.tenant, organization_unit=self.hq, code="SYN-R-UNLINKED",
            title="Same unit but no matching RiskControl", owner=self.admin
        )
        unrelated_treatment = RiskTreatment.objects.create(
            risk=unrelated_risk, strategy=RiskTreatment.Strategy.REDUCE, owner=self.admin
        )
        response = self.make_decision(register, risk_treatment=str(unrelated_treatment.id))
        self.assertEqual(response.status_code, 400)
        self.assertEqual(SoAControlDecision.objects.count(), 0)

    def test_readiness_is_read_only_and_never_issues_an_approved_soa(self):
        register = self.make_register()
        included = self.make_decision(register)
        self.assertEqual(included.status_code, 201, included.content)
        url = f"/api/v1/soa/registers/{register['id']}/readiness/"
        self.auth(self.viewer)
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200, response.content)
        payload = response.json()
        self.assertFalse(payload["formal_attestation_allowed"])
        self.assertFalse(payload["is_formal_statement_of_applicability"])
        self.assertEqual(payload["approval_workflow"], "not_available")
        self.assertEqual(payload["source_control_inventory_coverage"], "not_verified")
        self.assertEqual(payload["counts"]["included"], 1)
        self.assertEqual(payload["counts"]["decisions_with_findings"], 1)
        codes = payload["findings"][0]["codes"]
        self.assertIn("no_current_scoped_evidence_link", codes)
        self.assertIn("risk_treatment_not_approved", codes)
        self.assertGreaterEqual(len(payload["blocking_external_prerequisites"]), 3)
        self.assertEqual(SoARegister.objects.get(pk=register["id"]).status, "draft")
        self.assertEqual(self.client.post(url, {}, format="json").status_code, 405)

    def test_readiness_does_not_treat_expired_or_foreign_unit_evidence_as_scoped(self):
        register = self.make_register()
        included = self.make_decision(register)
        self.assertEqual(included.status_code, 201, included.content)
        decision_id = included.json()["id"]
        today = timezone.now()
        expired = Evidence.objects.create(
            tenant=self.tenant, organization_unit=self.hq, title="Expired synthetic evidence",
            owner=self.admin, valid_until=today - timedelta(days=1),
        )
        foreign_unit = Evidence.objects.create(
            tenant=self.tenant, organization_unit=self.ops, title="Wrong unit evidence",
            owner=self.admin, valid_until=today + timedelta(days=1),
        )
        for evidence in (expired, foreign_unit):
            EvidenceLink.objects.create(
                tenant=self.tenant, evidence=evidence, object_type="soa_decision",
                object_id=decision_id, relation_type="supports",
            )
        self.auth(self.admin)
        url = f"/api/v1/soa/registers/{register['id']}/readiness/"
        wrong = self.client.get(url)
        self.assertEqual(wrong.status_code, 200, wrong.content)
        self.assertIn("no_current_scoped_evidence_link", wrong.json()["findings"][0]["codes"])

        valid = Evidence.objects.create(
            tenant=self.tenant, organization_unit=self.hq, title="Current synthetic evidence",
            owner=self.admin, valid_until=today + timedelta(days=1),
        )
        EvidenceLink.objects.create(
            tenant=self.tenant, evidence=valid, object_type="soa_decision",
            object_id=decision_id, relation_type="supports",
        )
        correct = self.client.get(url)
        self.assertEqual(correct.status_code, 200, correct.content)
        codes = correct.json()["findings"][0]["codes"]
        self.assertNotIn("no_current_scoped_evidence_link", codes)
        self.assertIn("risk_treatment_not_approved", codes)
        self.assertEqual(correct.json()["evidence_sufficiency"], "requires_human_review")
        self.assertFalse(correct.json()["formal_attestation_allowed"])

    def test_readiness_does_not_leak_unassigned_scope_or_another_tenant(self):
        good = self.make_register(self.hq)
        forbidden = self.make_register(self.ops)
        self.auth(self.viewer)
        self.assertEqual(
            self.client.get(f"/api/v1/soa/registers/{good['id']}/readiness/").status_code, 200
        )
        self.assertEqual(
            self.client.get(f"/api/v1/soa/registers/{forbidden['id']}/readiness/").status_code, 404
        )
        self.auth(self.admin, tenant=self.foreign_tenant)
        foreign = self.client.get(f"/api/v1/soa/registers/{good['id']}/readiness/")
        self.assertIn(foreign.status_code, (403, 404))
        self.assertNotIn("Unapproved synthetic SoA decisions", foreign.content.decode())

    def test_legacy_word_export_is_draft_and_never_leaks_foreign_control_mapping(self):
        requirement = Requirement.objects.create(
            framework_version=self.version, code="SYN-1", title="Tenant-only synthetic requirement"
        )
        ControlRequirement.objects.create(
            tenant=self.tenant, requirement=requirement, control=self.control,
            mapping_type=ControlRequirement.MappingType.RELATED,
        )
        # Deliberately corrupted legacy mapping to the same framework version
        # from a foreign tenant; filtering by version alone would leak it.
        ControlRequirement.objects.create(
            tenant=self.foreign_tenant, requirement=requirement,
            control=self.foreign_control,
            mapping_type=ControlRequirement.MappingType.RELATED,
        )
        assessment = Assessment.objects.create(
            tenant=self.tenant, organization_unit=self.hq, framework_version=self.version,
            title="Synthetic assessment", owner=self.admin,
        )
        self.auth(self.admin)
        response = self.client.get(f"/api/v1/reports/soa/{assessment.id}/")
        self.assertEqual(response.status_code, 200, response.content)
        self.assertIn("soa-draft", response["Content-Disposition"])
        doc = Docx(BytesIO(response.content))
        paragraph_text = " ".join(p.text for p in doc.paragraphs)
        self.assertIn("NOT APPROVED", paragraph_text)
        self.assertIn("NOT AN APPROVED SoA", paragraph_text)
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(len(doc.tables[0].rows), 2)
        values = " ".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
        self.assertIn("Undetermined", values)
        self.assertIn("Not verified", values)
        self.assertIn(self.control.code, values)
        self.assertNotIn(self.foreign_control.code, values)
        self.assertNotIn(self.foreign_control.title, values)
        self.auth(self.viewer)
        self.assertEqual(
            self.client.get(f"/api/v1/reports/soa/{assessment.id}/").status_code,
            403,
        )

    def test_patch_can_update_decision_but_cannot_rebind_source_or_control(self):
        register = self.make_register()
        first = self.make_decision(register)
        self.assertEqual(first.status_code, 201, first.content)
        decision_id = first.json()["id"]
        self.auth(self.admin)
        path = f"/api/v1/soa/registers/{register['id']}/decisions/{decision_id}/"
        update = self.client.patch(path, {
            "decision": "undecided",
            "rationale": "",
            "control_implementation": None, "risk_treatment": None,
        }, format="json")
        self.assertEqual(update.status_code, 200, update.content)
        self.assertEqual(update.json()["decision"], "undecided")
        invalid = self.client.patch(path, {"control": str(self.second_control.id)}, format="json")
        self.assertEqual(invalid.status_code, 400)
        self.auth(self.viewer)
        self.assertEqual(self.client.patch(path, {"rationale": "no permission"}, format="json").status_code, 403)
