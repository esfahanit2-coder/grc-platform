from io import BytesIO

from django.db.models import Q
from django.http import HttpResponse
from docx import Document as Docx


def soa_docx(assessment):
    """Legacy draft cross-reference, never a signed Statement of Applicability."""
    from apps.controls.models import ControlRequirement, ControlImplementation

    doc = Docx()
    doc.add_heading(f"DRAFT applicability cross-reference — NOT APPROVED — {assessment.title}", 0)
    doc.add_paragraph(
        "DRAFT / NOT AN APPROVED SoA. A framework mapping is not a control-selection decision, "
        "and this export does not establish applicability, evidence sufficiency or certification."
    )
    table = doc.add_table(rows=1, cols=6)
    headers = [
        "Control", "Mapped requirement", "Applicability: UNDETERMINED",
        "Implementation (reported)", "Effectiveness (reported)", "Evidence: NOT VERIFIED",
    ]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header

    # Tenant ownership must be enforced on BOTH mapping and underlying control.
    # Merely matching a framework version can otherwise expose another tenant's
    # control title or a foreign mapping to a shared/global framework.
    maps = (
        ControlRequirement.objects.filter(
            requirement__framework_version=assessment.framework_version,
            deleted_at__isnull=True,
            control__deleted_at__isnull=True,
            requirement__deleted_at__isnull=True,
        )
        .filter(Q(tenant=assessment.tenant) | Q(tenant__isnull=True))
        .filter(Q(control__tenant=assessment.tenant) | Q(control__tenant__isnull=True))
        .select_related("control", "requirement")
        .order_by("control__code", "requirement__code")
    )
    for mapping in maps:
        implementation = (
            ControlImplementation.objects.filter(
                tenant=assessment.tenant, control=mapping.control,
                organization_unit=assessment.organization_unit, deleted_at__isnull=True,
            ).first()
            if assessment.organization_unit_id else None
        )
        cells = table.add_row().cells
        cells[0].text = f"{mapping.control.code} {mapping.control.title}"
        cells[1].text = mapping.requirement.code
        cells[2].text = "Undetermined — human control decision required"
        cells[3].text = implementation.implementation_status if implementation else "not_recorded"
        cells[4].text = implementation.effectiveness if implementation else "not_assessed"
        cells[5].text = "Not verified"

    output = BytesIO()
    doc.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    response["Content-Disposition"] = f'attachment; filename="soa-draft-{assessment.id}.docx"'
    return response

def document_docx(version):
    doc=Docx();doc.add_heading(version.document.title,0);doc.add_paragraph(version.content or '');b=BytesIO();doc.save(b);resp=HttpResponse(b.getvalue(),content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document');resp['Content-Disposition']=f'attachment; filename="{version.document.code}-{version.version}.docx"';return resp
